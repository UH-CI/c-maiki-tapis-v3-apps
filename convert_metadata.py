#!/usr/bin/env python3
"""Convert CMAIKI xlsx metadata to TSV format."""

import sys
import openpyxl

xlsx_file = sys.argv[1]
tsv_file = sys.argv[2]

# Template constants
HEADER_ROW = 11
DATA_START_ROW = 12

wb = openpyxl.load_workbook(xlsx_file, read_only=True, data_only=True)
ws = wb.active

header_row = list(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))[0]

# Find sample_id and samp_name columns
sample_id_col = None
samp_name_col = None

for idx, val in enumerate(header_row, 1):
    if val and str(val).lower() == 'sample_id':
        sample_id_col = idx
    elif val and str(val).lower() == 'samp_name':
        samp_name_col = idx

if not samp_name_col:
    print(f"ERROR: Could not find 'samp_name' column in row {HEADER_ROW}", file=sys.stderr)
    sys.exit(1)

# Write TSV
with open(tsv_file, 'w') as f:
    # Write header
    header_out = []
    for idx, val in enumerate(header_row, 1):
        if idx == sample_id_col:
            continue
        if idx == samp_name_col:
            header_out.append('ID')
        else:
            header_out.append(str(val) if val is not None else '')
    f.write('\t'.join(header_out) + '\n')
    
    # Write data rows
    for row in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
        samp_name_val = row[samp_name_col - 1] if len(row) >= samp_name_col else None
        if not samp_name_val:
            break
        
        data_out = []
        for idx, val in enumerate(row, 1):
            if idx == sample_id_col:
                continue
            data_out.append(str(val) if val is not None else '')
        
        f.write('\t'.join(data_out) + '\n')

wb.close()
print(f"Converted {xlsx_file} to {tsv_file}")